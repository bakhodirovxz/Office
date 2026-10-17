"""Generate the WP10 locale test workbook. Run: python -I gen_locale_xlsx.py OUT.xlsx
Column A = label, B = value/formula, C = number format (informational)."""
import sys, datetime
from openpyxl import Workbook

CASES = [
    ("general_big",   3500000.5,  "General"),
    ("general_frac",  0.1234567,  "General"),
    ("zero_dot_zero", 3500000,    "0.0"),
    ("grp_1dp",       3500000.5,  "#,##0.0"),
    ("grp_2dp",       3500000.5,  "#,##0.00"),
    ("pct",           0.255,      "0.00%"),
    ("sci",           12345.678,  "0.00E+00"),
    ("date_iso",      datetime.date(2026, 10, 12), "yyyy-mm-dd"),
    ("date_dmy",      datetime.date(2026, 10, 12), "dd.mm.yyyy"),
    ("date_long",     datetime.date(2026, 10, 12), "d mmmm yyyy"),
    ("date_dddd",     datetime.date(2026, 10, 12), "dddd, d mmm yyyy"),
    ("time_ampm",     datetime.time(15, 4, 5),     "h:mm:ss AM/PM"),
    ("tag419_num",    3500000.5,  "[$-419]#,##0.00"),
    ("tag443_num",    3500000.5,  "[$-443]#,##0.00"),
    ("tag419_date",   datetime.date(2026, 10, 12), "[$-419]d mmmm yyyy"),
    ("tag443_date",   datetime.date(2026, 10, 12), "[$-443]d mmmm yyyy"),
    ("tag_ru_cur",    3500000.5,  "[$\u20bd-419] #,##0.00"),
    ("tag_uz_cur",    3500000.5,  '#,##0.00 "so\'m"'),
    ("accounting",    3500000.5,  '_-* #,##0.00_-;-* #,##0.00_-;_-* "-"??_-;_-@_-'),
    ("f_text_num",    '=TEXT(3500000.5,"#,##0.00")', "General"),
    ("f_text_date",   '=TEXT(DATE(2026,10,12),"dd mmmm yyyy")', "General"),
    ("f_concat_num",  '=3500000.5&""', "General"),
    ("f_concat_date", '=DATE(2026,10,12)&""', "General"),
    ("f_sum_list",    '=SUM(1,2)', "General"),
]

def main(out):
    wb = Workbook(); ws = wb.active; ws.title = "L"
    ws.column_dimensions["A"].width = 16; ws.column_dimensions["B"].width = 30
    for i, (name, val, fmt) in enumerate(CASES, start=1):
        ws.cell(i, 1, name)
        c = ws.cell(i, 2, val); c.number_format = fmt
        ws.cell(i, 3, "'" + fmt if fmt.startswith("=") else fmt)
    wb.save(out)
    print("cases:", len(CASES))

if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "locale_cases.xlsx")
