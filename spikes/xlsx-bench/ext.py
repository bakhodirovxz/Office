"""Builds E:/office-target/xlsx-bench/ext/{source.xlsx,linked.xlsx}.
linked.xlsx is hand-assembled (zipfile) per ECMA-376 Part 1 18.14 / 18.17: workbook.xml <externalReferences>,
xl/externalLinks/externalLinkN.xml (sheetNames + cached sheetDataSet), rels with TargetMode=External.
Formulas use Excel's stored form [N]Sheet!Ref. externalLink1 = same-folder 'source.xlsx'; externalLink2 = absolute file:/// path.
usage: ext.py build | ext.py setsource   (setsource changes source values: A1..A5 x10, Data!B2=70)"""
import sys, zipfile, os
D = "E:/office-target/xlsx-bench/ext"; os.makedirs(D, exist_ok=True)
ABS = r"E:\office-target\xlsx-bench\ext\source.xlsx"
def source(mult):
    from openpyxl import Workbook
    wb = Workbook(); s = wb.active; s.title = "Sheet1"
    for i in range(1,6): s.cell(i,1,10*i*mult)
    d = wb.create_sheet("Data"); d["B2"] = 7*mult
    wb.save(f"{D}/source.xlsx")
NS = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
RNS = 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
def extlink(rid):
    cells1 = "".join(f'<row r="{i}"><cell r="A{i}"><v>{10*i}</v></cell></row>' for i in range(1,6))
    return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<externalLink {NS} {RNS}><externalBook r:id="{rid}">'
      '<sheetNames><sheetName val="Sheet1"/><sheetName val="Data"/></sheetNames><sheetDataSet>'
      f'<sheetData sheetId="0">{cells1}</sheetData><sheetData sheetId="1"><row r="2"><cell r="B2"><v>7</v></cell></row></sheetData>'
      '</sheetDataSet></externalBook></externalLink>')
def linked():
    cells = [("A1","[1]Sheet1!A1",10),("A2","[1]Sheet1!A2+1",21),("A3","SUM([1]Sheet1!A1:A5)",150),("A4","[1]Data!B2",7),
             ("B1","[2]Sheet1!A1",10),("B2","[2]Data!B2*2",14),("B3","A1+B1",20)]
    rows = {}
    for r,f,v in cells: rows.setdefault(int(r[1:]),[]).append(f'<c r="{r}"><f>{f}</f><v>{v}</v></c>')
    sd = "".join(f'<row r="{k}">{"".join(v)}</row>' for k,v in sorted(rows.items()))
    R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    files = {
     "[Content_Types].xml": '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/><Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/><Override PartName="/xl/externalLinks/externalLink1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.externalLink+xml"/><Override PartName="/xl/externalLinks/externalLink2.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.externalLink+xml"/></Types>',
     "_rels/.rels": f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="{R}/officeDocument" Target="xl/workbook.xml"/></Relationships>',
     "xl/workbook.xml": f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<workbook {NS} {RNS}><sheets><sheet name="Report" sheetId="1" r:id="rId1"/></sheets><externalReferences><externalReference r:id="rId2"/><externalReference r:id="rId3"/></externalReferences><calcPr calcId="191029"/></workbook>',
     "xl/_rels/workbook.xml.rels": f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="{R}/worksheet" Target="worksheets/sheet1.xml"/><Relationship Id="rId2" Type="{R}/externalLink" Target="externalLinks/externalLink1.xml"/><Relationship Id="rId3" Type="{R}/externalLink" Target="externalLinks/externalLink2.xml"/></Relationships>',
     "xl/worksheets/sheet1.xml": f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<worksheet {NS}><sheetData>{sd}</sheetData></worksheet>',
     "xl/externalLinks/externalLink1.xml": extlink("rId1"),
     "xl/externalLinks/_rels/externalLink1.xml.rels": f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="{R}/externalLinkPath" Target="source.xlsx" TargetMode="External"/></Relationships>',
     "xl/externalLinks/externalLink2.xml": extlink("rId1"),
     "xl/externalLinks/_rels/externalLink2.xml.rels": f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="{R}/externalLinkPath" Target="file:///{ABS}" TargetMode="External"/></Relationships>',
    }
    with zipfile.ZipFile(f"{D}/linked.xlsx","w",zipfile.ZIP_DEFLATED) as z:
        for k,v in files.items(): z.writestr(k,v)
if sys.argv[1]=="build": source(1); linked()
else: source(10)
