#!/usr/bin/env python3
"""Turn a review batch CSV (`cargo xt i18n review-export`) into the owner-facing XLSX.

Same layout as the glossary batch 0 workbook: columns A to H are the review columns (key, context, en, ru, uz-Latn draft,
uz-Cyrl draft, status, owner comment), the rest are helpers (scope, where-used, pass, en_hash, glossary terms) and are
ignored on import. The owner saves the sheet as CSV UTF-8 and runs `cargo xt i18n review-import`.

  python i18n/tools/review_xlsx.py i18n/review/batch-01.csv docs/phase1/i18n/review-batch-01.xlsx "Batch 1: menus, ribbon, files"
"""
import csv
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


def main(src: str, dst: str, title: str) -> None:
    with open(src, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f))
    head, body = rows[0], rows[1:]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(head)
    for r in body:
        ws.append(r)
    navy = PatternFill("solid", fgColor="1F3864")
    draft = PatternFill("solid", fgColor="FFF2CC")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = navy
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        row[4].fill = draft  # uz-Latn draft
        row[5].fill = draft  # uz-Cyrl draft
    widths = [14, 12, 34, 34, 40, 40, 18, 28, 11, 28, 8, 10, 26]
    for i, w in enumerate(widths[: len(head)], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.dimensions
    dv = DataValidation(type="list", formula1='"draft-unreviewed,reviewed,rejected"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"G2:G{len(body) + 1}")
    how = wb.create_sheet("How to review")
    for line in [
        f"{title.upper()} - DRAFT. Every Uzbek value in this file is an unreviewed proposal. Nothing is final until you mark it.",
        "",
        "For each row: read English (C) and Russian (D, what the upstream app ships). Check uz-Latn (E) and uz-Cyrl (F).",
        "Draft is right: set status (G) to reviewed. You prefer other text: type it over E and/or F and set G to reviewed.",
        "Wrong and no replacement: set G to rejected and say why in H. Not sure yet: leave G as draft-unreviewed.",
        "",
        "Context (B) tells apart two senses of one English word. 'where' (J) is the source file and line where the string appears.",
        "'glossary terms' (M) lists the batch 0 glossary terms the string uses: if you changed a glossary term, the string follows it",
        "(cargo xt i18n glossary-apply).",
        "Apostrophes: oʻ gʻ use U+02BB, the tutuq belgisi (maʼlumot) uses U+02BC. ASCII ' is reported as an error on import.",
        "Plural: Uzbek keeps the noun singular after a number ('5 ta fayl'), so one wording serves every count.",
        "Strings in ribbon and menu contexts that are much longer than the English are listed in docs/phase1/i18n/lint-report.md.",
        "",
        "Return it: File > Save As > CSV UTF-8 (Comma delimited), then   cargo xt i18n review-import <file.csv>",
    ]:
        how.append([line])
    how.column_dimensions["A"].width = 130
    how["A1"].font = Font(bold=True, color="C00000")
    Path(dst).parent.mkdir(parents=True, exist_ok=True)
    wb.save(dst)
    print(f"{dst}: {len(body)} rows")


if __name__ == "__main__":
    if len(sys.argv) != 4:
        print(__doc__)
        sys.exit(2)
    main(*sys.argv[1:])
