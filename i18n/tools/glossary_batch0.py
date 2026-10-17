#!/usr/bin/env python3
"""Glossary batch 0: the single source for i18n/glossary.tsv and the owner-facing
docs/phase1/i18n/glossary-batch0.{csv,xlsx}.

Every Uzbek value below is a DRAFT proposal. Nothing here is trusted until the owner marks it reviewed
(cargo xt i18n review-import). Orthography: official Uzbek Latin, o-reversed-comma U+02BB (oʻ, gʻ) and
U+02BC (tutuq belgisi, maʼlumot). Uzbek Cyrillic is written by hand, not transliterated by a program.

Run from the repo root after `cargo xt i18n extract` (it reads i18n/*.tsv for the ru reference and the
frequency columns):   python i18n/tools/glossary_batch0.py
Needs openpyxl for the .xlsx (the .csv and glossary.tsv are written either way).
"""
import csv
import hashlib
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
DATE = "2026-10-17"
STATUS = f"draft:agent-sonnet:{DATE}"

# (English, context, category, uz-Latn, uz-Cyrl, alt uz-Latn, alt uz-Cyrl, note)
T = [
    # ---- tabs and menus
    ("File", "", "menu", "Fayl", "Файл", "", "", "Established computing loan."),
    ("Edit", "", "menu", "Tahrirlash", "Таҳрирлаш", "Tahrir; Oʻzgartirish", "Таҳрир; Ўзгартириш", "Menu/tab name."),
    ("View", "", "menu", "Koʻrinish", "Кўриниш", "Koʻrish", "Кўриш", ""),
    ("Insert", "", "menu", "Kiritish", "Киритиш", "Qoʻshish; Joylashtirish", "Қўшиш; Жойлаштириш", "Tab/command that puts an object into the document. Add uses Qoʻshish, Paste uses Qoʻyish, so the three stay apart."),
    ("Format", "", "menu", "Format", "Формат", "Formatlash", "Форматлаш", "Noun and menu name; the verb form is the alternative."),
    ("Home", "", "menu", "Bosh sahifa", "Бош саҳифа", "Asosiy", "Асосий", "Ribbon tab, not a web home page."),
    ("Draw", "", "menu", "Chizish", "Чизиш", "Chizma", "Чизма", ""),
    ("Design", "", "menu", "Dizayn", "Дизайн", "Loyihalash", "Лойиҳалаш", ""),
    ("Layout", "", "menu", "Maket", "Макет", "Sahifa tuzilishi", "Саҳифа тузилиши", ""),
    ("Review", "", "menu", "Tekshirish", "Текшириш", "Retsenziya", "Рецензия", "Ribbon tab (comments, tracked changes)."),
    ("Help", "", "menu", "Yordam", "Ёрдам", "Qoʻllanma", "Қўлланма", ""),
    ("Tools", "", "menu", "Vositalar", "Воситалар", "Asboblar", "Асбоблар", ""),
    ("Options", "", "menu", "Parametrlar", "Параметрлар", "Sozlamalar", "Созламалар", "Options and Settings are separate words here so the two menus can differ."),
    ("Preferences", "", "menu", "Afzalliklar", "Афзалликлар", "Sozlamalar", "Созламалар", "PdfCraft menu. Owner to decide whether Preferences and Settings may share one word."),
    ("Tab", "ribbon", "menu", "Yorliq", "Ёрлиқ", "Tab", "Таб", "Ribbon tab (not a keyboard key)."),
    ("Window", "", "menu", "Oyna", "Ойна", "", "", ""),
    ("Slide Show", "", "menu", "Slaydlar namoyishi", "Слайдлар намойиши", "Namoyish", "Намойиш", "DeckCraft ribbon tab."),
    ("About", "", "menu", "Dastur haqida", "Дастур ҳақида", "Haqida", "Ҳақида", "About dialog; upstream product token stays inside the string (I5)."),
    ("Account", "", "menu", "Hisob", "Ҳисоб", "Akkaunt", "Аккаунт", ""),
    ("Recent", "", "menu", "Soʻnggi", "Сўнгги", "Yaqinda ishlatilgan", "Яқинда ишлатилган", "List of recently opened files."),
    # ---- commands
    ("New", "", "command", "Yangi", "Янги", "", "", ""),
    ("Open", "", "command", "Ochish", "Очиш", "", "", ""),
    ("Close", "", "command", "Yopish", "Ёпиш", "", "", ""),
    ("Save", "", "command", "Saqlash", "Сақлаш", "", "", ""),
    ("Save As", "", "command", "Boshqa nom bilan saqlash", "Бошқа ном билан сақлаш", "Nomlab saqlash", "Номлаб сақлаш", ""),
    ("Print", "", "command", "Chop etish", "Чоп этиш", "Bosib chiqarish", "Босиб чиқариш", ""),
    ("Print Preview", "", "command", "Chop etishdan oldin koʻrish", "Чоп этишдан олдин кўриш", "Oldindan koʻrish", "Олдиндан кўриш", ""),
    ("Export", "", "command", "Eksport qilish", "Экспорт қилиш", "Eksport", "Экспорт", ""),
    ("Import", "", "command", "Import qilish", "Импорт қилиш", "Import", "Импорт", ""),
    ("Share", "", "command", "Ulashish", "Улашиш", "Baham koʻrish", "Баҳам кўриш", ""),
    ("Undo", "", "command", "Qaytarish", "Қайтариш", "Bekor qilish", "Бекор қилиш", "Cancel takes Bekor qilish, so Undo avoids the same word."),
    ("Redo", "", "command", "Qayta bajarish", "Қайта бажариш", "Takrorlash", "Такрорлаш", ""),
    ("Cut", "", "command", "Kesib olish", "Кесиб олиш", "Qirqib olish", "Қирқиб олиш", "Clipboard cut. Crop (image) uses Qirqish."),
    ("Copy", "", "command", "Nusxalash", "Нусхалаш", "Nusxa olish", "Нусха олиш", ""),
    ("Paste", "", "command", "Qoʻyish", "Қўйиш", "Joylashtirish", "Жойлаштириш", ""),
    ("Delete", "", "command", "Oʻchirish", "Ўчириш", "Yoʻq qilish", "Йўқ қилиш", ""),
    ("Remove", "", "command", "Olib tashlash", "Олиб ташлаш", "Oʻchirish", "Ўчириш", "Takes something out of a list; the file is not deleted."),
    ("Add", "", "command", "Qoʻshish", "Қўшиш", "", "", ""),
    ("Select", "", "command", "Tanlash", "Танлаш", "Belgilash", "Белгилаш", ""),
    ("Select All", "", "command", "Hammasini tanlash", "Ҳаммасини танлаш", "Barchasini tanlash", "Барчасини танлаш", ""),
    ("Find", "", "command", "Topish", "Топиш", "Qidirish", "Қидириш", ""),
    ("Search", "", "command", "Qidirish", "Қидириш", "Izlash", "Излаш", ""),
    ("Replace", "", "command", "Almashtirish", "Алмаштириш", "Oʻrniga qoʻyish", "Ўрнига қўйиш", ""),
    ("Apply", "", "command", "Qoʻllash", "Қўллаш", "Tatbiq etish", "Татбиқ этиш", "Dialog button."),
    ("OK", "", "command", "OK", "OK", "", "", "Kept in Latin letters in both scripts, as in common Uzbek software."),
    ("Cancel", "", "command", "Bekor qilish", "Бекор қилиш", "", "", "Dialog button."),
    ("Yes", "", "command", "Ha", "Ҳа", "", "", ""),
    ("No", "", "command", "Yoʻq", "Йўқ", "", "", ""),
    ("Clear", "", "command", "Tozalash", "Тозалаш", "", "", ""),
    ("Browse", "", "command", "Koʻrib chiqish", "Кўриб чиқиш", "Tanlash", "Танлаш", "File chooser button."),
    ("Previous", "", "command", "Oldingi", "Олдинги", "", "", ""),
    ("Next", "", "command", "Keyingi", "Кейинги", "", "", ""),
    ("Back", "", "command", "Orqaga", "Орқага", "", "", ""),
    ("Zoom", "", "command", "Masshtab", "Масштаб", "Kattalashtirish", "Катталаштириш", ""),
    ("Rotate", "", "command", "Burish", "Буриш", "Aylantirish", "Айлантириш", ""),
    ("Sort", "", "command", "Saralash", "Саралаш", "Tartiblash", "Тартиблаш", ""),
    ("Merge", "", "command", "Birlashtirish", "Бирлаштириш", "", "", ""),
    ("Hide", "", "command", "Yashirish", "Яшириш", "", "", ""),
    ("Show", "", "command", "Koʻrsatish", "Кўрсатиш", "", "", ""),
    ("Move", "", "command", "Koʻchirish", "Кўчириш", "Surish", "Суриш", ""),
    ("Change", "", "command", "Oʻzgartirish", "Ўзгартириш", "", "", ""),
    ("Properties", "", "command", "Xossalar", "Хоссалар", "Xususiyatlar", "Хусусиятлар", ""),
    ("Link", "", "command", "Havola", "Ҳавола", "", "", ""),
    ("Align", "", "command", "Tekislash", "Текислаш", "", "", ""),
    ("Crop", "", "command", "Qirqish", "Қирқиш", "Kesish", "Кесиш", ""),
    ("Sign", "", "command", "Imzolash", "Имзолаш", "", "", ""),
    # ---- document objects
    ("Document", "", "object", "Hujjat", "Ҳужжат", "", "", ""),
    ("Page", "", "object", "Sahifa", "Саҳифа", "Bet", "Бет", ""),
    ("Table", "", "object", "Jadval", "Жадвал", "", "", ""),
    ("Cell", "", "object", "Katak", "Катак", "Yacheyka", "Ячейка", "Native word vs the Russian loan widely used in spreadsheets."),
    ("Row", "", "object", "Qator", "Қатор", "Satr", "Сатр", ""),
    ("Column", "", "object", "Ustun", "Устун", "", "", ""),
    ("Sheet", "", "object", "Varaq", "Варақ", "List", "Лист", "Russian loan list is common in speech."),
    ("Workbook", "", "object", "Ish kitobi", "Иш китоби", "Kitob", "Китоб", ""),
    ("Slide", "", "object", "Slayd", "Слайд", "Taqdimot varagʻi", "Тақдимот варағи", ""),
    ("Presentation", "", "object", "Taqdimot", "Тақдимот", "Prezentatsiya", "Презентация", ""),
    ("Text", "", "object", "Matn", "Матн", "", "", ""),
    ("Paragraph", "", "object", "Xatboshi", "Хатбоши", "Abzats", "Абзац", ""),
    ("Line", "text", "object", "Satr", "Сатр", "Qator", "Қатор", "A line of text."),
    ("Line", "shape", "object", "Chiziq", "Чизиқ", "", "", "A drawn line."),
    ("Word", "", "object", "Soʻz", "Сўз", "", "", ""),
    ("Character", "", "object", "Belgi", "Белги", "Harf", "Ҳарф", ""),
    ("Header", "", "object", "Yuqori kolontitul", "Юқори колонтитул", "", "", ""),
    ("Footer", "", "object", "Quyi kolontitul", "Қуйи колонтитул", "Pastki kolontitul", "Пастки колонтитул", ""),
    ("Heading", "", "object", "Sarlavha", "Сарлавҳа", "Boʻlim sarlavhasi", "Бўлим сарлавҳаси", "Heading style/level."),
    ("Title", "", "object", "Nomi", "Номи", "Sarlavha", "Сарлавҳа", "Document property."),
    ("Section", "", "object", "Boʻlim", "Бўлим", "", "", ""),
    ("Margins", "", "object", "Sahifa chetlari", "Саҳифа четлари", "Hoshiyalar", "Ҳошиялар", ""),
    ("Orientation", "", "object", "Yoʻnalish", "Йўналиш", "", "", ""),
    ("Portrait", "", "object", "Kitob yoʻnalishi", "Китоб йўналиши", "Tik", "Тик", ""),
    ("Landscape", "", "object", "Albom yoʻnalishi", "Албом йўналиши", "Yotiq", "Ётиқ", ""),
    ("Indent", "", "object", "Chekinish", "Чекиниш", "", "", ""),
    ("Spacing", "", "object", "Oraliq", "Оралиқ", "Interval", "Интервал", ""),
    ("Justify", "", "object", "Kenglik boʻyicha tekislash", "Кенглик бўйича текислаш", "Ikki tomonga tekislash", "Икки томонга текислаш", ""),
    ("Left", "", "object", "Chap", "Чап", "", "", ""),
    ("Right", "", "object", "Oʻng", "Ўнг", "", "", ""),
    ("Center", "", "object", "Markaz", "Марказ", "Markazga tekislash", "Марказга текислаш", ""),
    ("Top", "", "object", "Yuqori", "Юқори", "", "", ""),
    ("Bottom", "", "object", "Quyi", "Қуйи", "Pastki", "Пастки", ""),
    ("Bold", "", "format", "Qalin", "Қалин", "Yarim qalin", "Ярим қалин", ""),
    ("Italic", "", "format", "Kursiv", "Курсив", "Qiya", "Қия", ""),
    ("Underline", "", "format", "Tagiga chizish", "Тагига чизиш", "Tagchiziq", "Тагчизиқ", ""),
    ("Strikethrough", "", "format", "Ustidan chizish", "Устидан чизиш", "Ustiga chizilgan", "Устига чизилган", ""),
    ("Superscript", "", "format", "Yuqori indeks", "Юқори индекс", "", "", ""),
    ("Subscript", "", "format", "Quyi indeks", "Қуйи индекс", "Pastki indeks", "Пастки индекс", ""),
    ("Font", "", "format", "Shrift", "Шрифт", "", "", ""),
    ("Size", "", "format", "Oʻlcham", "Ўлчам", "", "", ""),
    ("Color", "", "format", "Rang", "Ранг", "", "", ""),
    ("Style", "", "format", "Uslub", "Услуб", "Stil", "Стил", ""),
    ("Theme", "", "format", "Tema", "Тема", "Mavzu", "Мавзу", "Subject (document property) takes Mavzu."),
    ("Border", "", "format", "Chegara", "Чегара", "Ramka", "Рамка", ""),
    ("Fill", "", "format", "Toʻldirish", "Тўлдириш", "Boʻyash", "Бўяш", ""),
    ("Picture", "", "object", "Rasm", "Расм", "Surat", "Сурат", ""),
    ("Shape", "", "object", "Shakl", "Шакл", "Figura", "Фигура", ""),
    ("Chart", "", "object", "Diagramma", "Диаграмма", "Grafik", "График", ""),
    ("Equation", "", "object", "Tenglama", "Тенглама", "Formula", "Формула", "Math equation in a document."),
    ("Formula", "", "object", "Formula", "Формула", "", "", "Spreadsheet formula."),
    ("Function", "", "object", "Funksiya", "Функция", "", "", ""),
    ("Value", "", "object", "Qiymat", "Қиймат", "", "", ""),
    ("Number", "", "object", "Son", "Сон", "Raqam", "Рақам", ""),
    ("Range", "", "object", "Diapazon", "Диапазон", "Oraliq", "Оралиқ", "Cell range."),
    ("Date", "", "object", "Sana", "Сана", "", "", ""),
    ("Comment", "", "object", "Izoh", "Изоҳ", "Sharh", "Шарҳ", ""),
    ("Annotation", "", "object", "Annotatsiya", "Аннотация", "Izohlash", "Изоҳлаш", "PDF annotation."),
    ("Signature", "", "object", "Imzo", "Имзо", "", "", ""),
    ("Certificate", "", "object", "Sertifikat", "Сертификат", "Guvohnoma", "Гувоҳнома", ""),
    ("Password", "", "object", "Parol", "Парол", "", "", ""),
    ("Security", "", "object", "Xavfsizlik", "Хавфсизлик", "", "", ""),
    ("Form", "", "object", "Forma", "Форма", "Anketa", "Анкета", "PDF form."),
    ("Field", "", "object", "Maydon", "Майдон", "", "", ""),
    ("Author", "", "object", "Muallif", "Муаллиф", "", "", ""),
    ("Version", "", "object", "Versiya", "Версия", "", "", ""),
    ("Folder", "", "object", "Jild", "Жилд", "Papka", "Папка", "Papka is the common spoken Russian loan."),
    ("Language", "", "settings", "Til", "Тил", "", "", ""),
    ("Spelling", "", "settings", "Imlo", "Имло", "", "", ""),
    ("Dark", "", "settings", "Qorongʻi", "Қоронғи", "Tungi", "Тунги", "Colour theme."),
    ("Light", "", "settings", "Yorugʻ", "Ёруғ", "Kunduzgi", "Кундузги", "Colour theme."),
    ("Auto", "", "settings", "Avto", "Авто", "Avtomatik", "Автоматик", ""),
    ("Default", "", "settings", "Standart", "Стандарт", "Odatiy", "Одатий", ""),
    ("Custom", "", "settings", "Maxsus", "Махсус", "", "", ""),
    ("None", "", "settings", "Yoʻq", "Йўқ", "Hech biri", "Ҳеч бири", ""),
    ("All", "", "settings", "Barchasi", "Барчаси", "Hammasi", "Ҳаммаси", ""),
    ("More", "", "settings", "Yana", "Яна", "Koʻproq", "Кўпроқ", ""),
    ("Name", "", "settings", "Nom", "Ном", "", "", ""),
    ("Type", "", "settings", "Tur", "Тур", "", "", ""),
    ("Location", "", "settings", "Joylashuv", "Жойлашув", "Manzil", "Манзил", ""),
    ("Width", "", "settings", "Eni", "Эни", "Kengligi", "Кенглиги", ""),
    ("Height", "", "settings", "Balandligi", "Баландлиги", "Boʻyi", "Бўйи", ""),
    ("Error", "", "settings", "Xato", "Хато", "", "", ""),
    ("Accessibility", "", "settings", "Maxsus imkoniyatlar", "Махсус имкониятлар", "Foydalanish imkoniyati", "Фойдаланиш имконияти", ""),
]

CATEGORY_ORDER = ["menu", "command", "object", "format", "settings"]


def sha1hex(s):
    return hashlib.sha1(s.encode("utf-8")).hexdigest()


def row_id(scope, ctx, src):
    return sha1hex(f"{scope}\x1f{ctx}\x1f{src}")[:12]


def esc(s):
    return s.replace("\\", "\\\\").replace("\n", "\\n").replace("\t", "\\t").replace("\r", "\\r")


def load_catalog():
    """Rows of the extracted catalog: (scope, ctx, src, ru, apps)."""
    rows = []
    for f in sorted((ROOT / "i18n").glob("*.tsv")):
        if f.stem == "glossary":
            continue
        for n, line in enumerate(f.read_text(encoding="utf-8").splitlines()):
            if n == 0 or not line:
                continue
            c = line.split("\t")
            if len(c) != 14:
                continue
            unesc = lambda s: re.sub(r"\\(.)", lambda m: {"n": "\n", "t": "\t", "r": "\r"}.get(m.group(1), m.group(1)), s)
            rows.append((f.stem, unesc(c[1]), unesc(c[2]), unesc(c[7]), c[11].split(",")))
    return rows


def checks():
    bad = []
    seen = set()
    for en, ctx, cat, lat, cyr, alat, acyr, note in T:
        if (en, ctx) in seen:
            bad.append(f"duplicate {en}/{ctx}")
        seen.add((en, ctx))
        for label, v in (("uz-Latn", lat), ("alt uz-Latn", alat)):
            for ch in "'\u2018\u2019`\u00b4":
                if ch in v:
                    bad.append(f"{en}: ASCII/typographic apostrophe in {label}: {v!r}")
        for label, v in (("uz-Cyrl", cyr), ("alt uz-Cyrl", acyr)):
            if re.search(r"[A-Za-z]", v) and v != "OK":
                bad.append(f"{en}: Latin letters in {label}: {v!r}")
        if re.search(r"[\u0400-\u04ff]", lat):
            bad.append(f"{en}: Cyrillic letters in uz-Latn")
        if cat not in CATEGORY_ORDER:
            bad.append(f"{en}: unknown category {cat}")
    return bad


def main():
    bad = checks()
    if bad:
        print("\n".join(bad))
        sys.exit(1)
    cat_rows = load_catalog()
    if not cat_rows:
        print("warning: i18n/*.tsv not found; run `cargo xt i18n extract` first (ru reference and frequency stay empty)")
    all_en = sorted({r[2] for r in cat_rows})
    lower = [s.lower() for s in all_en]

    def freq(term):
        pat = re.compile(r"(?<![a-z])" + re.escape(term.lower()) + r"(?![a-z])")
        return sum(1 for s in lower if pat.search(s))

    def ru_ref(term):
        best = ""
        for scope, ctx, src, ru, apps in cat_rows:
            if ctx == "" and src.rstrip("\u2026:") == term and ru:
                best = ru
                if src == term:
                    break
        return best

    def apps_exact(term):
        a = set()
        for scope, ctx, src, ru, apps in cat_rows:
            if ctx == "" and src.rstrip("\u2026:") == term:
                a.update(apps)
        return len(a)

    ordered = sorted(T, key=lambda t: (CATEGORY_ORDER.index(t[2]), t[0].lower(), t[1]))
    glossary_rows = []
    sheet = []
    for en, ctx, cat, lat, cyr, alat, acyr, note in ordered:
        rid = row_id("glossary", ctx, en)
        glossary_rows.append([
            rid, esc(ctx), esc(en), esc(lat), STATUS, esc(cyr), STATUS, esc(ru_ref(en)), "upstream" if ru_ref(en) else "todo", "", "P1",
            "wordcraft,gridcraft,pdfcraft,deckcraft", "", "; ".join(x for x in (f"category:{cat}", f"alt-latn:{alat}" if alat else "", f"alt-cyrl:{acyr}" if acyr else "", note) if x),
        ])
        sheet.append([
            rid, ctx, en, ru_ref(en), lat, cyr, "draft-unreviewed", "", cat, alat, acyr, note, freq(en), apps_exact(en),
            "glossary", "", "both", sha1hex(en)[:8],
        ])
    header = ["id", "context", "source_en", "uz_latn", "uz_latn_status", "uz_cyrl", "uz_cyrl_status", "ru", "ru_status", "ru_apps", "priority", "apps", "where", "note"]
    (ROOT / "i18n").mkdir(exist_ok=True)
    with open(ROOT / "i18n" / "glossary.tsv", "w", encoding="utf-8", newline="\n") as f:
        f.write("\t".join(header) + "\n")
        for r in sorted(glossary_rows, key=lambda r: r[0]):
            f.write("\t".join(r) + "\n")

    cols = ["key", "context", "en", "ru", "uz-Latn DRAFT", "uz-Cyrl DRAFT", "status", "owner comment", "category", "alt uz-Latn", "alt uz-Cyrl",
            "usage note", "strings containing term", "apps with exact string", "scope", "where", "pass", "en_hash"]
    out = ROOT / "docs" / "phase1" / "i18n"
    out.mkdir(parents=True, exist_ok=True)
    with open(out / "glossary-batch0.csv", "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, lineterminator="\r\n")
        w.writerow(cols)
        w.writerows(sheet)
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        print("openpyxl missing: wrote the .csv and glossary.tsv only")
        return
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Glossary batch 0"
    ws.append(cols)
    for r in sheet:
        ws.append(r)
    head = PatternFill("solid", fgColor="1F3864")
    draft = PatternFill("solid", fgColor="FFF2CC")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for idx in (4, 5):  # uz-Latn DRAFT, uz-Cyrl DRAFT
            row[idx].fill = draft
    widths = [14, 9, 24, 28, 28, 28, 18, 30, 10, 26, 26, 44, 12, 12, 10, 10, 8, 10]
    for i, wdt in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = wdt
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.dimensions
    dv = DataValidation(type="list", formula1='"draft-unreviewed,reviewed,rejected"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"G2:G{len(sheet) + 1}")
    ws.column_dimensions["A"].hidden = False
    rd = wb.create_sheet("How to review")
    for line in [
        "GLOSSARY BATCH 0 - DRAFT. Every Uzbek value in this file is an unreviewed proposal. Nothing is final until you mark it.",
        "",
        "What to do for each row:",
        "1. Read the English term (column C) and the Russian reference (column D, what the upstream app already ships).",
        "2. Check uz-Latn DRAFT (E) and uz-Cyrl DRAFT (F). Columns J and K show 1-2 alternatives where usage varies; column L says why.",
        "3. If the draft is right, set status (G) to: reviewed.",
        "4. If you prefer another term, type your term over the draft cell (E and/or F) and set status to: reviewed.",
        "   The tool keeps your text exactly as you typed it (it only checks the apostrophe rule and tells you if it fails).",
        "5. If the term is wrong and you have no replacement, set status to: rejected and write why in 'owner comment' (H).",
        "6. Leave status as draft-unreviewed for rows you want to look at later.",
        "",
        "Apostrophes: use oʻ gʻ with U+02BB, and the tutuq belgisi (maʼlumot) with U+02BC. ASCII ' is reported as an error.",
        "Columns A-H are the review columns. Columns I-R are helper columns and are ignored on import.",
        "Column M counts how many distinct English UI strings of the four apps contain the term; column N counts the apps that use the bare term.",
        "",
        "How to return it: File > Save As > CSV UTF-8 (Comma delimited), then:   cargo xt i18n review-import <file.csv>",
        "Terms with context 'text', 'shape' or 'ribbon' are the same English word used in different senses.",
    ]:
        rd.append([line])
    rd.column_dimensions["A"].width = 130
    rd["A1"].font = Font(bold=True, color="C00000")
    wb.save(out / "glossary-batch0.xlsx")
    print(f"glossary: {len(T)} terms; wrote glossary.tsv, glossary-batch0.csv, glossary-batch0.xlsx")


if __name__ == "__main__":
    main()
