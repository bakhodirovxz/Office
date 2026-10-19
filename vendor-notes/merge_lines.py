"""Merge line-per-paragraph fragments produced by pdf2docx back into real paragraphs."""
import copy, re, sys, docx
from docx.enum.text import WD_ALIGN_PARAGRAPH

src, dst = sys.argv[1], sys.argv[2]
d = docx.Document(src)
END = re.compile(r'[.;:!?…»”"]\s*$')

def continues(p, q):
    a, b = p.text.rstrip(), q.text.lstrip()
    if not a or not b or END.search(a):
        return False
    if p._p.getparent() is not q._p.getparent():
        return False
    if is_heading(p) or is_heading(q):
        return False
    # Next line continues the sentence: starts lower-case or with a digit/punctuation,
    # or the current line is a full-width body line cut mid-sentence.
    return b[0].islower() or b[0].isdigit() or b[0] in '(-–«“"' or len(a) >= 40

def is_heading(p):
    runs = [r for r in p.runs if r.text.strip()]
    return bool(runs) and all(r.bold for r in runs)

merged = 0
paras = d.paragraphs
i = 0
while i < len(paras) - 1:
    p, q = paras[i], paras[i + 1]
    if continues(p, q):
        if not p.text.endswith((' ', '-')):
            p.add_run(' ')
        for r in q._p.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}r'):
            p._p.append(copy.deepcopy(r))
        q._p.getparent().remove(q._p)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paras.pop(i + 1)
        merged += 1
    else:
        i += 1
# Collapse the double spaces left over from justified PDF lines.
for p in d.paragraphs:
    for r in p.runs:
        t = r.text
        if len(p.text) > 150 and '\t' in t:
            t = t.replace('\t', ' ')  # justification tabs inside running text
        if '  ' in t:
            t = re.sub(r' {2,}', ' ', t)
        if t != r.text:
            r.text = t
d.save(dst)
print('merged', merged, 'paragraphs now', len([p for p in d.paragraphs if p.text.strip()]))
